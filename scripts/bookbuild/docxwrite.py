"""Render parsed blocks (see mdparse.py) into a polished python-docx Document."""

import hashlib
import os
import posixpath

import requests
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BRAND_BLUE = RGBColor(0x25, 0x63, 0xEB)  # matches the site's brand-primary colour
GREY = RGBColor(0x55, 0x55, 0x55)
CALLOUT_FILL = "EAF2FF"


# ---------------------------------------------------------------------------
# Document setup
# ---------------------------------------------------------------------------

def new_document():
    doc = Document()

    section = doc.sections[0]
    section.page_height, section.page_width = Inches(11.69), Inches(8.27)  # A4
    for attr in ("top_margin", "bottom_margin"):
        setattr(section, attr, Inches(1))
    for attr in ("left_margin", "right_margin"):
        setattr(section, attr, Inches(1.1))

    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = "Georgia"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    for level, size in ((1, 24), (2, 18), (3, 14), (4, 12)):
        h = styles[f"Heading {level}"]
        h.font.name = "Georgia"
        h.font.bold = True
        h.font.size = Pt(size)
        h.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
        h.paragraph_format.space_before = Pt(18 if level <= 2 else 12)
        h.paragraph_format.space_after = Pt(8)

    if "Byline" not in [s.name for s in styles]:
        byline = styles.add_style("Byline", styles["Normal"].type)
        byline.base_style = normal
        byline.font.italic = True
        byline.font.size = Pt(10)
        byline.font.color.rgb = GREY

    if "Caption" not in [s.name for s in styles]:
        cap = styles.add_style("Caption", styles["Normal"].type)
        cap.base_style = normal
        cap.font.italic = True
        cap.font.size = Pt(9.5)
        cap.font.color.rgb = GREY

    return doc


def add_title_page(doc, manifest, cover_image_path, compiled_date):
    for _ in range(4):
        doc.add_paragraph()
    if cover_image_path and os.path.exists(cover_image_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(cover_image_path, width=Inches(1.6))
        doc.add_paragraph()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(manifest.get("title", "Untitled"))
    run.font.size = Pt(36)
    run.font.bold = True
    run.font.name = "Georgia"

    if manifest.get("subtitle"):
        sub = doc.add_paragraph()
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = sub.add_run(manifest["subtitle"])
        r.font.size = Pt(18)
        r.font.color.rgb = GREY
        r.font.name = "Georgia"

    for _ in range(6):
        doc.add_paragraph()

    foot = doc.add_paragraph()
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = foot.add_run(f"Compiled {compiled_date} from monkton-kent.uk")
    r.font.size = Pt(10)
    r.font.color.rgb = GREY
    r.italic = True

    doc.add_page_break()


def add_toc(doc):
    heading = doc.add_heading("Contents", level=1)
    heading.paragraph_format.page_break_before = False

    p = doc.add_paragraph()
    run = p.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-2" \\h \\z \\u'
    fld_char_sep = OxmlElement("w:fldChar")
    fld_char_sep.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Right-click and choose “Update Field” to generate the contents."
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")

    r_element = run._r
    r_element.append(fld_char_begin)
    r_element.append(instr)
    r_element.append(fld_char_sep)
    r_element.append(placeholder)
    r_element.append(fld_char_end)

    doc.add_page_break()


def add_header_footer(doc, title):
    section = doc.sections[-1]
    section.different_first_page_header_footer = True

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_page_number_field(fp)

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = hp.add_run(title)
    r.font.size = Pt(9)
    r.font.color.rgb = GREY
    r.italic = True


def _add_page_number_field(paragraph):
    run = paragraph.add_run()
    for kind, text in (("begin", None), ("instr", "PAGE"), ("end", None)):
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        run._r.append(el)
    run.font.size = Pt(9)
    run.font.color.rgb = GREY


def add_part_heading(doc, title, subtitle=None):
    doc.add_page_break()
    doc.add_heading(title, level=1)
    if subtitle:
        p = doc.add_paragraph()
        r = p.add_run(subtitle)
        r.italic = True
        r.font.color.rgb = GREY
        p.paragraph_format.space_after = Pt(16)


# ---------------------------------------------------------------------------
# Block rendering
# ---------------------------------------------------------------------------

def render_blocks(doc, blocks, *, level_offset, link_targets, current_dir, image_cache_dir,
                   chapter_bookmark=None, page_break_before_chapter=True, source_dir):
    first_heading = True
    for block in blocks:
        t = block["type"]

        if t == "heading":
            level = min(block["level"] + level_offset, 4)
            if first_heading and page_break_before_chapter:
                doc.add_page_break()
            h = doc.add_heading(level=level)
            if first_heading and chapter_bookmark:
                _start_bookmark(h, chapter_bookmark)
            add_runs(h, block["runs"], link_targets, current_dir)
            if first_heading and chapter_bookmark:
                _end_bookmark(h)
            first_heading = False

        elif t == "para":
            p = doc.add_paragraph()
            add_runs(p, block["runs"], link_targets, current_dir)
            _style_if_byline(p)

        elif t == "list":
            for item in block["items"]:
                if block["ordered"] and item["marker"]:
                    # explicit reference numbers (e.g. tablet/plot numbers,
                    # not necessarily sequential) - kept as literal text
                    # rather than handed to Word's auto-numbering.
                    p = doc.add_paragraph()
                    p.paragraph_format.left_indent = Inches(0.25)
                    r = p.add_run(item["marker"] + " ")
                    r.bold = True
                else:
                    p = doc.add_paragraph(style="List Bullet")
                add_runs(p, item["runs"], link_targets, current_dir)

        elif t == "table":
            _render_table(doc, block)

        elif t == "callout":
            _render_callout(doc, block, level_offset, link_targets, current_dir, image_cache_dir, source_dir)

        elif t == "quote":
            for inner in block["blocks"]:
                if inner["type"] == "para":
                    p = doc.add_paragraph()
                    p.paragraph_format.left_indent = Inches(0.4)
                    add_runs(p, inner["runs"], link_targets, current_dir, force_italic=True)

        elif t == "figure":
            _render_figure(doc, block, current_dir, image_cache_dir, source_dir)


def _style_if_byline(paragraph):
    text = paragraph.text.strip()
    if text.startswith("*") is False and len(paragraph.runs) == 1 and paragraph.runs[0].italic and not paragraph.runs[0].bold:
        # A lone italicised line (byline, e.g. "*Terry Marsh*") - style smaller/grey.
        for r in paragraph.runs:
            r.font.color.rgb = GREY
            r.font.size = Pt(10)
    if text.startswith("Originally published in the Monkton News"):
        for r in paragraph.runs:
            r.font.size = Pt(9.5)
            r.font.color.rgb = GREY


def _render_table(doc, block):
    n_cols = len(block["header"])
    table = doc.add_table(rows=1, cols=n_cols)
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    for cell, runs in zip(hdr_cells, block["header"]):
        p = cell.paragraphs[0]
        for run in runs:
            r = p.add_run(run["text"])
            r.bold = True
    for row in block["rows"]:
        cells = table.add_row().cells
        for cell, runs in zip(cells, row):
            p = cell.paragraphs[0]
            for run in runs:
                r = p.add_run(run["text"])
                r.bold = run["bold"]
                r.italic = run["italic"]
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def _render_callout(doc, block, level_offset, link_targets, current_dir, image_cache_dir, source_dir):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.rows[0].cells[0]
    _shade_cell(cell, CALLOUT_FILL)
    cell.paragraphs[0].text = ""
    title_p = cell.paragraphs[0]
    r = title_p.add_run(block["title"])
    r.bold = True
    r.font.color.rgb = BRAND_BLUE

    for inner in block["blocks"]:
        if inner["type"] == "para":
            p = cell.add_paragraph()
            add_runs(p, inner["runs"], link_targets, current_dir)
        elif inner["type"] == "list":
            for item in inner["items"]:
                p = cell.add_paragraph(style="List Bullet")
                add_runs(p, item["runs"], link_targets, current_dir)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def _shade_cell(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _render_figure(doc, block, current_dir, image_cache_dir, source_dir):
    local_path = resolve_image(block["src"], current_dir, image_cache_dir, source_dir)
    if local_path and os.path.exists(local_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        try:
            p.add_run().add_picture(local_path, width=Inches(5.5))
        except Exception:
            pass
    if block["caption_runs"]:
        cap = doc.add_paragraph(style="Caption")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_runs(cap, block["caption_runs"], {}, current_dir)


# ---------------------------------------------------------------------------
# Runs, links, bookmarks
# ---------------------------------------------------------------------------

def add_runs(paragraph, runs, link_targets, current_dir, force_italic=False):
    for run in runs:
        if not run["text"]:
            continue
        link = run["link"]
        if link:
            kind, target = resolve_link(link, current_dir, link_targets)
        else:
            kind, target = None, None

        if kind == "internal":
            _add_internal_hyperlink(paragraph, run["text"], target, run, force_italic)
        elif kind == "external":
            _add_external_hyperlink(paragraph, run["text"], target, run, force_italic)
        else:
            r = paragraph.add_run(run["text"])
            r.bold = run["bold"]
            r.italic = run["italic"] or force_italic
            if run["sup"]:
                r.font.superscript = True
            if run["code"]:
                r.font.name = "Consolas"


def resolve_link(url, current_dir, link_targets):
    if url.startswith(("http://", "https://", "mailto:")):
        return "external", url
    # normalise a relative markdown link to a source/-rooted posix path
    joined = posixpath.normpath(posixpath.join(current_dir, url.split("#")[0]))
    if joined in link_targets:
        return "internal", link_targets[joined]
    return "plain", None


def _add_external_hyperlink(paragraph, text, url, run, force_italic):
    part = paragraph.part
    r_id = part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = _styled_run_element(text, run, force_italic, color="2563EB", underline=True)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def _add_internal_hyperlink(paragraph, text, bookmark_name, run, force_italic):
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), bookmark_name)
    new_run = _styled_run_element(text, run, force_italic, color="2563EB", underline=True)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def _styled_run_element(text, run, force_italic, color=None, underline=False):
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    if run["bold"]:
        rpr.append(OxmlElement("w:b"))
    if run["italic"] or force_italic:
        rpr.append(OxmlElement("w:i"))
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rpr.append(u)
    if color:
        c = OxmlElement("w:color")
        c.set(qn("w:val"), color)
        rpr.append(c)
    if run["sup"]:
        vert = OxmlElement("w:vertAlign")
        vert.set(qn("w:val"), "superscript")
        rpr.append(vert)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    return r


_bookmark_id = [0]


def _start_bookmark(paragraph, name):
    _bookmark_id[0] += 1
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(_bookmark_id[0]))
    start.set(qn("w:name"), name)
    # w:pPr, if present, must stay the first child of w:p - insert just after it.
    p_el = paragraph._p
    ppr = p_el.find(qn("w:pPr"))
    insert_at = 1 if ppr is not None else 0
    p_el.insert(insert_at, start)
    paragraph._bookmark_id = _bookmark_id[0]


def _end_bookmark(paragraph):
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(getattr(paragraph, "_bookmark_id", _bookmark_id[0])))
    paragraph._p.append(end)


# ---------------------------------------------------------------------------
# Images
# ---------------------------------------------------------------------------

def resolve_image(src, current_dir, image_cache_dir, source_dir):
    if src.startswith(("http://", "https://")):
        return download_cached(src, image_cache_dir)
    path = os.path.normpath(os.path.join(source_dir, current_dir, src))
    return path


def download_cached(url, image_cache_dir):
    os.makedirs(image_cache_dir, exist_ok=True)
    ext = os.path.splitext(url.split("?")[0])[1] or ".img"
    digest = hashlib.sha256(url.encode()).hexdigest()[:24]
    cache_path = os.path.join(image_cache_dir, digest + ext)
    if os.path.exists(cache_path):
        return cache_path
    try:
        resp = requests.get(url, timeout=20)
        resp.raise_for_status()
        with open(cache_path, "wb") as f:
            f.write(resp.content)
        return cache_path
    except Exception as exc:
        print(f"  ! image download failed, skipping: {url} ({exc})")
        return None
